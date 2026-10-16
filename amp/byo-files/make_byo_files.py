#!/usr/bin/env python
"""Generates the three synthetic sample files for the Bring-your-own lab.

    python make_byo_files.py            # writes loan_book.csv, rate_sheet.xlsx, complaint_log.pdf here

Everything is fictional (Northwind Auto Finance). loan_book.csv uses a fixed seed, so the
output is identical every run. The loan math and the credit tiers match the
northwind_finance_tools package (the same formulas and the same table).

Needs only the standard library for loan_book.csv; rate_sheet.xlsx needs openpyxl and
complaint_log.pdf needs reportlab (a scratch environment is enough: these are not lab dependencies).
"""
from __future__ import annotations

import csv
import random
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
SEED = 20261104
AS_OF = date(2026, 9, 30)
ROWS = 200

# Same table as northwind_finance_tools.tools.CREDIT_TIERS (illustrative, synthetic).
TIERS = [
    ("Super prime", 781, 850, 4.5, 6.5),
    ("Prime", 661, 780, 6.5, 9.5),
    ("Near prime", 601, 660, 9.5, 14.0),
    ("Subprime", 501, 600, 14.0, 19.5),
    ("Deep subprime", 300, 500, 19.5, 24.0),
]
TERM_ADJUSTMENTS = [(36, -0.50), (48, -0.25), (60, 0.00), (72, 0.50)]
STATES = ["TX", "GA", "FL", "OH", "AZ", "NC", "TN", "CA", "MI", "PA", "IL", "CO"]
FIRST_NAMES = ["Alex", "Blake", "Casey", "Drew", "Emery", "Finley", "Gray", "Harper", "Indy", "Jules",
               "Kai", "Lane", "Morgan", "Noel", "Oakley", "Parker", "Quinn", "Reese", "Sage", "Tatum"]


def _cents(x: float) -> float:
    return round(x + 1e-9, 2)


def payment(principal: float, apr: float, months: int) -> float:
    r = apr / 1200
    if r == 0:
        return _cents(principal / months)
    f = (1 + r) ** months
    return _cents(principal * r * f / (f - 1))


def balance_after(principal: float, apr: float, months: int, paid: int) -> float:
    r = apr / 1200
    pay = payment(principal, apr, months)
    bal = principal
    for _ in range(paid):
        bal = _cents(bal - (pay - _cents(bal * r)))
    return max(bal, 0.0)


def months_between(a: date, b: date) -> int:
    return (b.year - a.year) * 12 + (b.month - a.month)


def tier_for(score: int):
    return next(t for t in TIERS if t[1] <= score <= t[2])


def make_loan_book(path: Path) -> None:
    rng = random.Random(SEED)
    odd = random.Random(SEED + 1)  # separate stream for the mispriced loans
    fields = ["account_id", "origination_date", "state", "dealer_id", "vehicle_condition", "credit_score",
              "credit_tier", "original_principal", "apr", "term_months", "monthly_payment", "payments_made",
              "current_balance", "days_past_due", "status", "autopay"]
    rows = []
    for i in range(ROWS):
        score = min(850, max(300, int(rng.gauss(675, 85))))
        tier = tier_for(score)
        term = rng.choices([36, 48, 60, 72], weights=[1, 2, 5, 3])[0]
        adj = dict(TERM_ADJUSTMENTS)[term]
        apr = round(rng.uniform(tier[3], tier[4]) + adj, 1)
        apr = max(apr, 0.0)
        if odd.random() < 0.06:  # a few loans priced outside their tier's band (for the pricing-check example)
            apr = round(max(apr + odd.choice([-1, 1]) * odd.uniform(2.0, 3.5), 0.5), 1)
        if rng.random() < 0.04:  # a few promotional 0 percent loans
            apr = 0.0
        condition = rng.choices(["new", "used"], weights=[4, 6])[0]
        principal = float(round(rng.uniform(9000, 31000 if condition == "used" else 52000) / 50) * 50)
        orig = date(2021, 10, 1).fromordinal(date(2021, 10, 1).toordinal() + rng.randint(0, 1800))
        elapsed = max(months_between(orig, AS_OF), 0)
        missed = rng.choices([0, 1, 2, 3, 4], weights=[82, 8, 5, 3, 2])[0] if score < 700 else rng.choices([0, 1], weights=[96, 4])[0]
        made = max(min(elapsed - missed, term), 0)
        pay = payment(principal, apr, term)
        bal = balance_after(principal, apr, term, made)
        paid_off = made >= term
        missed = 0 if paid_off else min(missed, elapsed)
        dpd = 0 if missed == 0 else 30 * missed + rng.randint(0, 9)
        if paid_off:
            status = "paid_off"
        elif dpd == 0:
            status = "current"
        elif dpd < 60:
            status = "30_days_late"
        elif dpd < 90:
            status = "60_days_late"
        else:
            status = "90_plus_days_late"
        rows.append({
            "account_id": f"LB-{10001 + i}",
            "origination_date": orig.isoformat(),
            "state": rng.choice(STATES),
            "dealer_id": f"DLR-{rng.randint(1, 24):03d}",
            "vehicle_condition": condition,
            "credit_score": score,
            "credit_tier": tier[0],
            "original_principal": f"{principal:.2f}",
            "apr": f"{apr:.1f}",
            "term_months": term,
            "monthly_payment": f"{pay:.2f}",
            "payments_made": made,
            "current_balance": f"{bal:.2f}",
            "days_past_due": dpd,
            "status": status,
            "autopay": "Y" if rng.random() < 0.55 else "N",
        })
    rows.sort(key=lambda r: r["account_id"])
    with path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def make_rate_sheet(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="161616")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col, w in zip("ABCDEFG", widths):
            ws.column_dimensions[col].width = w

    ws = wb.active
    ws.title = "Rate bands"
    sheet(ws, ["Credit tier", "Min score", "Max score", "APR low % (60 mo)", "APR high % (60 mo)"],
          [list(t) for t in TIERS], [18, 12, 12, 20, 20])
    ws2 = wb.create_sheet("Term adjustments")
    sheet(ws2, ["Term (months)", "APR adjustment (points)"], [list(t) for t in TERM_ADJUSTMENTS], [16, 26])
    ws3 = wb.create_sheet("Notes")
    for line in [
        "Northwind Auto Finance - illustrative rate sheet (SYNTHETIC)",
        "Effective date: 2026-10-01 (fictional).",
        "This sheet is a teaching example. It is not a real rate sheet and not an offer of credit.",
        "Final APR = a rate inside the tier's band for 60 months, plus the term adjustment.",
        "The band is the same table the credit_tier_lookup tool uses.",
        "Example: score 712 is Prime (6.5% to 9.5% at 60 months); a 72-month loan adds 0.50 points.",
    ]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 100
    ws3["A1"].font = Font(bold=True)
    wb.save(path)


COMPLAINTS = [
    ("CMP-2026-0101", "2026-08-03", "Phone", "LB-10011", "Alex R.",
     "Customer says the autopay draft came out twice in July and left the checking account overdrawn. "
     "Asks for a refund of the duplicate payment and the overdraft fee."),
    ("CMP-2026-0102", "2026-08-11", "Email", "LB-10013", "Blake T.",
     "Customer paid off the loan on 2026-07-14 and has not received the lien release or title. "
     "States it has been almost four weeks and the dealer needs the title to complete a trade-in."),
    ("CMP-2026-0103", "2026-08-19", "Web form", "LB-10024", "Casey M.",
     "Customer received two collection calls in one day after being told by a representative that "
     "the account was current. Requests that the calls stop until the balance is confirmed."),
    ("CMP-2026-0104", "2026-08-27", "Letter", "LB-10113", "Drew P.",
     "Customer believes a late fee was charged during a declared disaster period and wants it "
     "reversed. Says they will contact their state regulator if it is not resolved."),
    ("CMP-2026-0105", "2026-09-02", "Phone", "LB-10131", "Emery S.",
     "Customer asked for a payoff quote for a trade-in and was told it would take a week. "
     "Dealer needs it in two days. Customer asks for the quote to be expedited."),
    ("CMP-2026-0106", "2026-09-09", "Email", "LB-10156", "Finley K.",
     "Customer says the monthly statement shows an interest rate different from the contract. "
     "Attached the contract page. Asks for a corrected statement and an explanation."),
    ("CMP-2026-0107", "2026-09-15", "Phone", "LB-10170", "Gray L.",
     "Active-duty servicemember says they were not offered the rate reduction they requested "
     "after receiving orders. Asks to speak with a supervisor."),
    ("CMP-2026-0108", "2026-09-22", "Web form", "LB-10199", "Harper D.",
     "Customer reports an unfamiliar charge from a third-party insurance add-on they did not "
     "agree to at signing. Asks for the add-on to be cancelled and the charge refunded."),
]


def make_complaint_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter, title="Northwind Auto Finance - complaint log (synthetic)",
                            author="Northwind Auto Finance (fictional)")
    story = [
        Paragraph("Northwind Auto Finance - Customer Complaint Log", styles["Title"]),
        Paragraph("Synthetic sample for the workshop. Fictional lender, fictional customers. "
                  "Period: 2026-08-01 to 2026-09-30.", styles["Normal"]),
        Spacer(1, 14),
    ]
    for cid, day, channel, acct, name, text in COMPLAINTS:
        story.append(Paragraph(f"<b>{cid}</b> - received {day} via {channel}", styles["Heading3"]))
        story.append(Paragraph(f"Account: {acct} | Customer: {name}", styles["Normal"]))
        story.append(Paragraph(text, styles["BodyText"]))
        story.append(Spacer(1, 8))
    doc.build(story)


if __name__ == "__main__":
    make_loan_book(HERE / "loan_book.csv")
    make_rate_sheet(HERE / "rate_sheet.xlsx")
    make_complaint_pdf(HERE / "complaint_log.pdf")
    print("wrote loan_book.csv, rate_sheet.xlsx, complaint_log.pdf in", HERE, file=sys.stderr)
